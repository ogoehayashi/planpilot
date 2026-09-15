"""Strict compact JSON/Excel factory adapter. All quantities use integer base units."""
from copy import deepcopy
from datetime import datetime, timedelta, timezone
import json
from pathlib import Path

from .models import FactoryData, Order, Operation

LOCAL_TIMEZONE = timezone(timedelta(hours=8))


def integer(value, name, minimum=0):
    if type(value) is not int or value < minimum or value > 10_000_000:
        raise ValueError(f"{name} must be an integer between {minimum} and 10000000")
    return value


def identifier(value, name):
    if not isinstance(value, str) or not value.strip() or len(value) > 128:
        raise ValueError(f"{name} must be a non-empty string of at most 128 characters")
    return value


def read_factory(path):
    path = Path(path)
    if path.suffix.lower() == ".json":
        return json.loads(path.read_text(encoding="utf-8-sig"))
    if path.suffix.lower() == ".xlsx":
        return _excel_bundle(path)
    raise ValueError("supported factory formats are .json and .xlsx")


def load_factory(path):
    return factory_from_dict(read_factory(path))


def factory_from_dict(raw):
    if not isinstance(raw, dict) or not isinstance(raw.get("orders"), list):
        raise ValueError("orders array is required")
    if raw.get("allow_optional_lot_splitting") is True:
        raise ValueError("optional lot splitting is unsupported")
    if len(raw["orders"]) > 100 or sum(len(o.get("operations", [])) for o in raw["orders"]) > 200:
        raise ValueError("compact adapter supports at most 100 orders and 200 operations")
    origin = datetime.fromisoformat(raw.get("planning_start", "2026-01-01T00:00:00+08:00"))
    if origin.tzinfo is None:
        raise ValueError("planning_start must include a timezone")
    horizon = integer(raw.get("horizon", 7200), "horizon", 1)
    orders, ids = [], set()
    for row in raw["orders"]:
        oid = identifier(row["order_id"], "order_id")
        if oid in ids:
            raise ValueError("duplicate order_id")
        ids.add(oid)
        if not isinstance(row.get("operations"), list) or not row["operations"]:
            raise ValueError("each order requires operations")
        operations = []
        for i, op in enumerate(row["operations"]):
            worker = identifier(op["worker_id"], "worker_id") if op.get("worker_id") is not None else None
            material = identifier(op["material_id"], "material_id") if op.get("material_id") is not None else None
            quantity = integer(op.get("material_qty", 0), "material_qty")
            if bool(material) != bool(quantity):
                raise ValueError("material_id and positive material_qty must occur together")
            operations.append(Operation(i + 1, identifier(op["machine_id"], "machine_id"), integer(op["duration"], "duration", 1),
                                        worker, material, quantity, integer(op.get("release", 0), "release"),
                                        integer(op.get("changeover", 0), "changeover"), op.get("required_skill")))
        due = row.get("due_at")
        if due is None:
            date = datetime.fromisoformat(str(row["due_date"]))
            if date.tzinfo is None:
                date = date.replace(tzinfo=LOCAL_TIMEZONE)
            due = int((date - origin).total_seconds() // 60)
        orders.append(Order(oid, identifier(row["product_id"], "product_id"), integer(row["quantity"], "quantity", 1),
                            integer(due, "due_at"), tuple(operations), integer(row.get("priority", 0), "priority")))
    inventory = {}
    for material, value in raw.get("inventory", {}).items():
        identifier(material, "material_id")
        if isinstance(value, int):
            value = {"quantity": value}
        batches, batch_ids = [], set()
        for row in value.get("batches", [dict(value, batch_id=material)]):
            bid = identifier(row["batch_id"], "batch_id")
            if bid in batch_ids:
                raise ValueError("duplicate inventory batch_id")
            batch_ids.add(bid)
            batches.append({"batch_id": bid, "quantity": integer(row["quantity"], "quantity"),
                            "available_at": integer(row.get("available_at", 0), "available_at")})
        inventory[material] = {"batches": batches}
    workers, machines = deepcopy(raw.get("workers", [])), deepcopy(raw.get("machines", []))
    for rows, field in ((workers, "worker_id"), (machines, "machine_id")):
        resource_ids = [identifier(r[field], field) for r in rows]
        if len(resource_ids) != len(set(resource_ids)):
            raise ValueError(f"duplicate {field}")
    worker_ids = {w["worker_id"] for w in workers}
    for worker in workers:
        skills = worker.get("skills", [])
        if not isinstance(skills, list):
            raise ValueError("worker skills must be an array")
        seen_skills = set()
        for skill in skills:
            if isinstance(skill, str):
                identifier(skill, "skill")
                continue
            sid = identifier(skill["skill_id"], "skill_id")
            if sid in seen_skills or not 1 <= integer(skill["proficiency_level"], "proficiency_level", 1) <= 5 or type(skill.get("is_primary")) is not bool:
                raise ValueError("invalid or duplicate worker skill")
            seen_skills.add(sid)
    machine_ids = {m["machine_id"] for m in machines} or {op.machine_id for o in orders for op in o.operations}
    for order in orders:
        for op in order.operations:
            if op.machine_id not in machine_ids or (op.worker_id and op.worker_id not in worker_ids):
                raise ValueError("operation references unknown resource")
    shifts, maintenance = deepcopy(raw.get("shifts", [])), deepcopy(raw.get("maintenance", []))
    for row in shifts + maintenance:
        start, end = integer(row["start"], "start"), integer(row["end"], "end", 1)
        if end <= start:
            raise ValueError("resource window must have end > start")
        if not row.get("machine_id") and not row.get("worker_id"):
            raise ValueError("resource window requires machine_id or worker_id")
        if row.get("machine_id") and row["machine_id"] not in machine_ids or row.get("worker_id") and row["worker_id"] not in worker_ids:
            raise ValueError("resource window references unknown resource")
        if row in shifts and row.get("window_type", "REGULAR") != "REGULAR":
            raise ValueError("compact adapter accepts regular shifts only; overtime requires the contract calendar adapter")
    return FactoryData(tuple(orders), inventory, tuple(maintenance), tuple(workers), tuple(shifts), tuple(machines), horizon)


def _excel_bundle(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if not {"Orders", "Operations", "Inventory"}.issubset(workbook.sheetnames):
            raise ValueError("workbook must contain Orders, Operations and Inventory sheets")
        def rows(name):
            if name not in workbook.sheetnames:
                return []
            values = iter(workbook[name].values)
            headers = next(values, ())
            if not headers or any(not isinstance(h, str) or not h for h in headers) or len(headers) != len(set(headers)):
                raise ValueError(f"invalid or duplicate headers in {name}")
            return [{k: v for k, v in zip(headers, row) if v is not None} for row in values if any(v is not None for v in row)]
        orders = rows("Orders")
        by_order = {r["order_id"]: [] for r in orders}
        for row in rows("Operations"):
            if row["order_id"] not in by_order:
                raise ValueError("operation references unknown order")
            row = dict(row)
            row["duration"] = row.get("duration", row.get("duration_min"))
            by_order[row["order_id"]].append(row)
        for order in orders:
            ops = by_order[order["order_id"]]
            if any("operation_no" not in op for op in ops):
                raise ValueError("Excel Operations requires operation_no")
            if len({op["operation_no"] for op in ops}) != len(ops):
                raise ValueError("duplicate operation_no")
            order["operations"] = sorted(ops, key=lambda op: op["operation_no"])
            if isinstance(order.get("due_date"), datetime):
                order["due_date"] = order["due_date"].isoformat()
        inventory = {}
        for row in rows("Inventory"):
            inventory.setdefault(row["material_id"], {"batches": []})["batches"].append(
                {"batch_id": row.get("batch_id", row.get("source_id")), "quantity": row.get("quantity", row.get("quantity_base_units")), "available_at": row.get("available_at", 0)})
        workers = rows("Workers")
        for worker in workers:
            worker["skills"] = json.loads(worker.get("skills", "[]"))
        settings = {r["key"]: r["value"] for r in rows("Settings")}
        return {**settings, "orders": orders, "inventory": inventory, "workers": workers,
                "machines": rows("Machines"), "shifts": rows("Shifts"), "maintenance": rows("Maintenance")}
    finally:
        workbook.close()
