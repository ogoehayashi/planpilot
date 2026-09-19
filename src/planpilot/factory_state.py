"""Production V1.8 workbook migration, validation and immutable state storage.

The compact adapter in :mod:`planpilot.domain.importer` remains a solver bridge.
This module owns the 17-sheet workbook boundary and never repairs bad records or
invents missing planning values.
"""
from __future__ import annotations

from copy import deepcopy
from datetime import date, datetime, time, timedelta, timezone
import hashlib
import json
from pathlib import Path
from typing import Any

from .persistence import canonical
from .validation import validate_tool_payload


SHEETS = (
    "README", "Assumptions", "Machines", "Workers", "Products", "Routing",
    "Inventory", "Orders", "Events", "Baseline Schedule", "Objectives",
    "Approval Policy", "Evaluation Cases", "Plan Output Schema",
    "Shift Calendar", "Worker Skills", "Changeovers",
)
REQUIRED_COLUMNS = {
    "README": {"as_of_time", "dataset_version", "generator_seed"},
    "Assumptions": {"key", "value"},
    "Machines": {"machine_id", "machine_group_id", "capacity_minutes"},
    "Workers": {"worker_id", "name", "available"},
    "Products": {"product_id", "requires_material", "max_lot_size", "material_id", "material_uom", "quantity_per_finished_unit_base_units"},
    "Routing": {"product_id", "operation_no", "operation_type", "machine_id", "required_skill", "duration_min"},
    "Inventory": {"material_id", "source_id", "source_type", "confirmed", "quantity_base_units", "available_at", "material_uom"},
    "Orders": {"order_id", "product_id", "quantity", "due_date", "priority"},
    "Events": {"event_id", "event_type", "Payload_JSON"},
    "Baseline Schedule": {"record_status"},
    "Objectives": {"profile", "delivery_weight", "overtime_weight", "changeover_weight", "stability_weight"},
    "Approval Policy": {"action", "decision", "approver"},
    "Evaluation Cases": {"case_id", "test_type", "scenario", "pass_condition"},
    "Plan Output Schema": {"field", "required"},
    "Shift Calendar": {"calendar_window_id", "window_type", "start_at", "end_at", "resource_type", "resource_id", "overtime_allowed"},
    "Worker Skills": {"worker_id", "skill_id", "proficiency_level", "is_primary"},
    "Changeovers": {"machine_group_id", "from_product_id", "to_product_id", "changeover_minutes"},
}
SGT = timezone(timedelta(hours=8))


def _issue(code: str, entity_type: str, entity_id: str | None, field: str | None, message: str, severity: str = "ERROR") -> dict:
    return {"code": code, "severity": severity, "entity_type": entity_type,
            "entity_id": entity_id, "field": field, "message": message}


def _identifier(value: Any) -> bool:
    return isinstance(value, str) and bool(value.strip()) and len(value) <= 128


def _integer(value: Any, minimum: int = 0) -> bool:
    return type(value) is int and minimum <= value <= 10_000_000


def _timestamp(value: Any) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.min, SGT)
    elif isinstance(value, str):
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    else:
        raise ValueError("timestamp must be an ISO-8601 string")
    if parsed.tzinfo is None or parsed.utcoffset() != timedelta(hours=8):
        raise ValueError("timestamp must use Asia/Singapore UTC+08:00")
    return parsed


def _minutes(value: Any, origin: datetime) -> int:
    if type(value) is int:
        return value
    return int((_timestamp(value) - origin).total_seconds() // 60)


def _json_safe(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, dict):
        return {key: _json_safe(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_json_safe(item) for item in value]
    return value


def _read(path: Path) -> tuple[dict[str, list[dict]], list[dict], str]:
    from openpyxl import load_workbook
    raw_bytes = path.read_bytes()
    workbook = load_workbook(path, read_only=True, data_only=True)
    structural: list[dict] = []
    tables: dict[str, list[dict]] = {}
    try:
        missing = [name for name in SHEETS if name not in workbook.sheetnames]
        for name in missing:
            structural.append(_issue("REQUIRED_FIELD_MISSING", "workbook", None, name, f"required sheet {name} is missing"))
        for name in SHEETS:
            if name not in workbook.sheetnames:
                tables[name] = []
                continue
            values = iter(workbook[name].values)
            headers = list(next(values, ()))
            if not headers or any(not _identifier(h) for h in headers) or len(headers) != len(set(headers)):
                structural.append(_issue("INVALID_VALUE", "sheet", name, "headers", "headers must be unique non-empty strings"))
                tables[name] = []
                continue
            absent = sorted(REQUIRED_COLUMNS[name] - set(headers))
            for field in absent:
                structural.append(_issue("REQUIRED_FIELD_MISSING", "sheet", name, field, f"required column {field} is missing"))
            tables[name] = [
                {key: value for key, value in zip(headers, row) if value is not None}
                for row in values if any(value is not None for value in row)
            ]
    finally:
        workbook.close()
    return _json_safe(tables), structural, hashlib.sha256(raw_bytes).hexdigest()


def normalize_workbook(path: str | Path, as_of_time: str | None = None) -> dict:
    """Read all 17 sheets and produce a validated solver state plus evidence."""
    path = Path(path)
    tables, structural, source_sha256 = _read(path)
    errors, warnings, quarantine = list(structural), [], []
    if structural:
        validation = {"status": "INVALID", "errors": errors, "warnings": warnings,
                      "quarantined_entities": quarantine}
        return {"state": {}, "validation": validation, "dataset_version": "unknown",
                "entity_counts": {name: len(rows) for name, rows in tables.items()},
                "warnings": [], "source_sha256": source_sha256, "source_tables": tables}

    readme = tables["README"]
    if len(readme) != 1:
        errors.append(_issue("INVALID_VALUE", "README", None, None, "README must contain exactly one data row"))
        origin = datetime(1970, 1, 1, tzinfo=SGT)
        dataset_version = "unknown"
    else:
        dataset_version = str(readme[0].get("dataset_version", ""))
        if not dataset_version or not _integer(readme[0].get("generator_seed"), 0):
            errors.append(_issue("INVALID_VALUE", "README", None, "dataset_version", "dataset_version and integer generator_seed are required"))
        try:
            origin = _timestamp(as_of_time or readme[0].get("as_of_time"))
        except (TypeError, ValueError) as exc:
            errors.append(_issue("INVALID_VALUE", "README", None, "as_of_time", str(exc)))
            origin = datetime(1970, 1, 1, tzinfo=SGT)

    bad: dict[str, set[str]] = {name: set() for name in ("machine", "worker", "product", "routing", "inventory", "order", "event", "calendar", "skill", "changeover")}

    def reject(code, kind, entity_id, field, message):
        value = str(entity_id) if entity_id is not None else None
        issue = _issue(code, kind, value, field, message)
        errors.append(issue); quarantine.append(issue)
        if kind in bad and value is not None:
            bad[kind].add(value)

    assumptions = {}
    for row in tables["Assumptions"]:
        key = row.get("key")
        if not _identifier(key) or key in assumptions:
            reject("INVALID_VALUE", "assumption", key, "key", "assumption keys must be unique identifiers")
        else:
            assumptions[key] = row.get("value")
    required_assumptions = {"overtime_cap_hours", "max_overtime_min_per_worker_per_day", "changeover_reference_min", "stability_drift_min", "allow_optional_lot_splitting"}
    for key in sorted(required_assumptions - set(assumptions)):
        reject("REQUIRED_FIELD_MISSING", "assumption", key, "value", "required assumption is missing; no default was substituted")
    if assumptions.get("allow_optional_lot_splitting") is not False:
        reject("OPTIONAL_LOT_SPLITTING_UNSUPPORTED", "assumption", "allow_optional_lot_splitting", "value", "optional lot splitting must be false")
    for key in ("overtime_cap_hours", "max_overtime_min_per_worker_per_day", "changeover_reference_min", "stability_drift_min"):
        if key in assumptions and not _integer(assumptions[key], 1):
            reject("INVALID_VALUE", "assumption", key, "value", "assumption must be a positive integer")

    objective_profiles = {"Balanced", "Delivery First", "Cost First"}
    seen_profiles = set()
    for row in tables["Objectives"]:
        profile = row.get("profile")
        weights = [row.get(key) for key in ("delivery_weight", "overtime_weight", "changeover_weight", "stability_weight")]
        if profile not in objective_profiles or profile in seen_profiles or any(type(value) not in (int, float) or not 0 <= value <= 1 for value in weights) or abs(sum(weights) - 1) > 1e-9:
            reject("INVALID_VALUE", "objective", profile, None, "profile must be unique and its four weights must be numbers in [0,1] summing to 1")
        else:
            seen_profiles.add(profile)
    for profile in sorted(objective_profiles - seen_profiles):
        reject("REQUIRED_FIELD_MISSING", "objective", profile, None, "required objective profile is missing")

    approval_actions = {"generate_or_simulate_plan", "assign_qualified_secondary_skill", "add_overtime", "publish_plan"}
    approval_decisions = {"AUTO_ALLOW", "REQUIRE_CONFIRMATION", "REQUIRE_APPROVAL"}
    seen_actions = set()
    for row in tables["Approval Policy"]:
        action = row.get("action")
        if action not in approval_actions or action in seen_actions or row.get("decision") not in approval_decisions:
            reject("INVALID_VALUE", "approval_policy", action, None, "approval action must be unique and use a controlled decision")
        else:
            seen_actions.add(action)
    for action in sorted(approval_actions - seen_actions):
        reject("REQUIRED_FIELD_MISSING", "approval_policy", action, None, "required approval action is missing")

    eval_ids = {row.get("case_id") for row in tables["Evaluation Cases"]}
    expected_evals = {f"EVAL-{number:03d}" for number in range(1, 31)}
    if eval_ids != expected_evals or any(not _identifier(row.get("scenario")) or not _identifier(row.get("test_type")) or not isinstance(row.get("pass_condition"), str) or not row["pass_condition"] for row in tables["Evaluation Cases"]):
        reject("INVALID_VALUE", "evaluation_cases", None, None, "Evaluation Cases must contain complete EVAL-001 through EVAL-030 rows")
    output_fields = {row.get("field"): row.get("required") for row in tables["Plan Output Schema"]}
    if output_fields != {"operations": True, "kpis": True}:
        reject("INVALID_VALUE", "plan_output_schema", None, None, "operations and kpis must both be required")
    if tables["Baseline Schedule"] != [{"record_status": "none"}]:
        reject("INVALID_VALUE", "baseline_schedule", None, "record_status", "empty demo baseline must be represented by one 'none' row")

    machines, machine_ids, groups = [], set(), {}
    for row in tables["Machines"]:
        mid = row.get("machine_id")
        if not _identifier(mid) or mid in machine_ids or not _identifier(row.get("machine_group_id")) or not _integer(row.get("capacity_minutes"), 1):
            reject("MACHINE_CAPACITY_INVALID", "machine", mid, None, "machine id/group must be valid and capacity_minutes a positive integer")
            continue
        machine_ids.add(mid); groups.setdefault(row["machine_group_id"], []).append(mid); machines.append(dict(row))

    workers, worker_ids = [], set()
    for row in tables["Workers"]:
        wid = row.get("worker_id")
        if not _identifier(wid) or wid in worker_ids or type(row.get("available")) is not bool:
            reject("WORKER_AVAILABILITY_INVALID", "worker", wid, None, "worker id must be unique and available must be boolean")
            continue
        worker_ids.add(wid); workers.append({**row, "skills": []})

    skill_workers: dict[str, list[dict]] = {}
    seen_skills = set()
    for row in tables["Worker Skills"]:
        wid, sid = row.get("worker_id"), row.get("skill_id")
        key = f"{wid}:{sid}"
        if wid not in worker_ids or not _identifier(sid) or key in seen_skills or not _integer(row.get("proficiency_level"), 1) or row.get("proficiency_level", 0) > 5 or type(row.get("is_primary")) is not bool:
            reject("BROKEN_REFERENCE" if wid not in worker_ids else "INVALID_VALUE", "skill", key, None, "skill must reference a worker and declare proficiency 1..5 plus is_primary")
            continue
        seen_skills.add(key)
        clean = {"skill_id": sid, "proficiency_level": row["proficiency_level"], "is_primary": row["is_primary"]}
        next(w for w in workers if w["worker_id"] == wid)["skills"].append(clean)
        skill_workers.setdefault(sid, []).append({"worker_id": wid, **clean})

    products, product_ids = [], set()
    for row in tables["Products"]:
        pid = row.get("product_id")
        valid = _identifier(pid) and pid not in product_ids and type(row.get("requires_material")) is bool and _integer(row.get("max_lot_size"), 1)
        if row.get("requires_material"):
            valid = valid and _identifier(row.get("material_id")) and row.get("material_uom") in {"EA", "G", "ML"} and _integer(row.get("quantity_per_finished_unit_base_units"), 1)
        if not valid:
            reject("BOM_ROW_MISSING" if row.get("requires_material") else "INVALID_VALUE", "product", pid, None, "product requires a positive lot size and complete integer-base-unit BOM")
            continue
        product_ids.add(pid)
        bom = [] if not row["requires_material"] else [{"material_id": row["material_id"], "material_uom": row["material_uom"], "quantity_per_unit": row["quantity_per_finished_unit_base_units"]}]
        products.append({"product_id": pid, "max_lot_size": row["max_lot_size"], "requires_material": row["requires_material"], "bom": bom})

    route_by_product: dict[str, list[dict]] = {}
    seen_route = set()
    for row in tables["Routing"]:
        pid, no = row.get("product_id"), row.get("operation_no")
        rid = f"{pid}:{no}"
        code = None
        if pid not in product_ids or row.get("machine_id") not in machine_ids or row.get("required_skill") not in skill_workers:
            code = "BROKEN_REFERENCE"
        elif not _integer(no, 1) or (pid, no) in seen_route or row.get("operation_type") not in {"PRODUCTION", "INSPECTION"} or not _integer(row.get("duration_min"), 1):
            code = "ROUTING_INVALID"
        if code:
            reject(code, "routing", rid, None, "routing must have unique sequence, controlled operation type, positive duration and valid resource/skill references")
            bad["product"].add(str(pid)); continue
        candidates = sorted(skill_workers[row["required_skill"]], key=lambda item: (not item["is_primary"], -item["proficiency_level"], item["worker_id"]))
        seen_route.add((pid, no))
        route_by_product.setdefault(pid, []).append({"operation_no": no, "operation_type": row["operation_type"], "machine_id": row["machine_id"], "worker_id": candidates[0]["worker_id"], "duration": row["duration_min"], "required_skill": row["required_skill"]})
    for pid in product_ids:
        route = sorted(route_by_product.get(pid, []), key=lambda item: item["operation_no"])
        if not route or route[-1]["operation_type"] != "INSPECTION" or [r["operation_no"] for r in route] != list(range(1, len(route) + 1)):
            reject("ROUTING_INVALID", "routing", pid, None, "product routing must be contiguous and end with INSPECTION")
            bad["product"].add(pid)

    inventory: dict[str, dict] = {}
    source_ids = set()
    for row in tables["Inventory"]:
        sid, material = row.get("source_id"), row.get("material_id")
        valid = _identifier(sid) and sid not in source_ids and _identifier(material) and row.get("material_uom") in {"EA", "G", "ML"} and _integer(row.get("quantity_base_units")) and row.get("source_type") in {"ON_HAND", "CONFIRMED_INBOUND"} and type(row.get("confirmed")) is bool
        try:
            available = _minutes(row.get("available_at", 0), origin)
            valid = valid and available >= 0
        except (TypeError, ValueError):
            valid = False; available = 0
        if not valid:
            reject("MATERIAL_UOM_INVALID" if row.get("material_uom") not in {"EA", "G", "ML"} else "INVALID_VALUE", "inventory", sid, None, "inventory bucket must use integer base units, a controlled UOM/source type and a valid availability time")
            continue
        source_ids.add(sid)
        entry = inventory.setdefault(material, {"material_uom": row["material_uom"], "batches": []})
        if entry["material_uom"] != row["material_uom"]:
            reject("MATERIAL_UOM_INVALID", "inventory", sid, "material_uom", "all buckets for one material must share a UOM"); continue
        entry["batches"].append({"batch_id": sid, "quantity": row["quantity_base_units"], "available_at": available, "source_type": row["source_type"]})

    clean_products = []
    for product in products:
        missing_material = any(b["material_id"] not in inventory or inventory[b["material_id"]]["material_uom"] != b["material_uom"] for b in product["bom"])
        if product["product_id"] in bad["product"] or missing_material:
            if missing_material:
                reject("QUARANTINE_DEPENDENCY", "product", product["product_id"], "bom", "product depends on missing or quarantined inventory")
            bad["product"].add(product["product_id"])
        else:
            clean_products.append(product)

    orders = []
    order_ids = set()
    for row in tables["Orders"]:
        oid, pid = row.get("order_id"), row.get("product_id")
        valid = _identifier(oid) and oid not in order_ids and _integer(row.get("quantity"), 1) and _integer(row.get("priority"), 0)
        try:
            due = _timestamp(row.get("due_date")) if not isinstance(row.get("due_date"), str) or "T" in row.get("due_date", "") else datetime.combine(date.fromisoformat(row["due_date"]), time.min, SGT)
            due_at = int((due - origin).total_seconds() // 60)
            valid = valid and 0 <= due_at <= 10_000_000
        except (TypeError, ValueError):
            valid = False; due_at = 0
        if not valid:
            reject("INVALID_VALUE", "order", oid, None, "order requires unique id, positive integer quantity, integer priority and an in-range date")
            continue
        order_ids.add(oid)
        if pid not in product_ids:
            reject("BROKEN_REFERENCE", "order", oid, "product_id", "order references an unknown product"); continue
        if pid in bad["product"]:
            reject("QUARANTINE_DEPENDENCY", "order", oid, "product_id", f"order depends on quarantined product {pid}"); continue
        orders.append({"order_id": oid, "product_id": pid, "quantity": row["quantity"], "due_date": due.date().isoformat(), "due_at": due_at, "priority": row["priority"], "operations": deepcopy(sorted(route_by_product[pid], key=lambda item: item["operation_no"]))})

    shifts, calendar_ids = [], set()
    for row in tables["Shift Calendar"]:
        cid = row.get("calendar_window_id")
        try:
            start, end = _minutes(row.get("start_at"), origin), _minutes(row.get("end_at"), origin)
        except (TypeError, ValueError):
            start = end = 0
        valid = _identifier(cid) and cid not in calendar_ids and row.get("window_type") in {"REGULAR", "OVERTIME"} and type(row.get("overtime_allowed")) is bool and (row["window_type"] == "OVERTIME") == row["overtime_allowed"] and 0 <= start < end
        resource_type, resource_id = row.get("resource_type"), row.get("resource_id")
        targets = groups.get(resource_id, []) if resource_type == "MACHINE_GROUP" else ([resource_id] if resource_type == "WORKER_GROUP" and resource_id in worker_ids else [])
        if not valid or not targets:
            reject("SHIFT_WINDOW_REFERENCE_INVALID", "calendar", cid, "resource_id", "calendar row must reference a declared resource and use a valid window"); continue
        calendar_ids.add(cid)
        for target in targets:
            shifts.append({"calendar_window_id": cid, "window_type": row["window_type"], "overtime_allowed": row["overtime_allowed"], "start": start, "end": end, "machine_id" if resource_type == "MACHINE_GROUP" else "worker_id": target})

    changeovers = []
    seen_changeovers = set()
    for row in tables["Changeovers"]:
        group, source, target = row.get("machine_group_id"), row.get("from_product_id"), row.get("to_product_id")
        key = (group, source, target)
        if group not in groups or source not in product_ids or target not in product_ids or key in seen_changeovers or not _integer(row.get("changeover_minutes")):
            reject("CHANGEOVER_TRANSITION_MISSING", "changeover", ":".join(map(str, key)), None, "changeover must be a unique complete transition with non-negative integer minutes"); continue
        seen_changeovers.add(key)
        for machine_id in groups[group]:
            changeovers.append({"machine_id": machine_id, "from_product_id": source, "to_product_id": target, "minutes": row["changeover_minutes"]})
    for group in groups:
        for source in product_ids:
            for target in product_ids:
                if (group, source, target) not in seen_changeovers:
                    reject("CHANGEOVER_TRANSITION_MISSING", "changeover", f"{group}:{source}:{target}", None, "required transition is missing")

    event_types = {"URGENT_ORDER", "MACHINE_BREAKDOWN", "MATERIAL_DELAY", "WORKER_ABSENCE", "PROMPT_INJECTION", "QUANTITY_REVISION", "DUE_DATE_PULL_IN"}
    events, event_ids = [], set()
    for row in tables["Events"]:
        eid = row.get("event_id")
        if not _identifier(eid) or eid in event_ids or row.get("event_type") not in event_types:
            reject("INVALID_VALUE", "event", eid, "event_type", "event id/type is invalid or duplicated"); continue
        event_ids.add(eid)
        reference_ok = True
        if row["event_type"] in {"URGENT_ORDER", "QUANTITY_REVISION", "DUE_DATE_PULL_IN"}:
            reference_ok = row.get("order_id") in order_ids
        elif row["event_type"] == "MACHINE_BREAKDOWN":
            reference_ok = row.get("machine_id") in machine_ids
        elif row["event_type"] == "MATERIAL_DELAY":
            reference_ok = row.get("material_id") in inventory
        elif row["event_type"] == "WORKER_ABSENCE":
            reference_ok = row.get("worker_id") in worker_ids
        if not reference_ok:
            reject("BROKEN_REFERENCE", "event", eid, None, "event references an unknown entity"); continue
        if row["event_type"] == "PROMPT_INJECTION":
            reject("PROMPT_INJECTION", "event", eid, "Payload_JSON", "untrusted instruction-like payload was quarantined")
            continue
        safe = {k: v for k, v in row.items() if k != "Payload_JSON"}
        events.append(safe)

    workdays = sorted({(origin + timedelta(minutes=row["start"])).date() for row in shifts})
    expected_days = [origin.date() + timedelta(days=index) for index in range(5)]
    if workdays[:5] != expected_days:
        errors.append(_issue("DATE_OUT_OF_RANGE", "workbook", None, "Shift Calendar", "five consecutive working days from as_of_time are required"))
        horizon = 0
    else:
        horizon_end = datetime.combine(expected_days[-1] + timedelta(days=1), time.min, SGT)
        horizon = int((horizon_end - origin).total_seconds() // 60)
    state = {
        "planning_start": origin.isoformat(timespec="seconds"), "horizon": horizon,
        "products": clean_products, "orders": orders, "inventory": inventory,
        "machines": machines, "workers": workers, "shifts": shifts,
        "maintenance": [], "changeovers": changeovers, "assumptions": assumptions,
        "events": events,
    }
    status = "INVALID" if any(issue["entity_type"] in {"workbook", "sheet", "README", "assumption"} for issue in errors) else ("VALID_WITH_QUARANTINE" if quarantine else "VALID")
    validation = {"status": status, "errors": errors, "warnings": warnings, "quarantined_entities": quarantine}
    return {"state": state, "validation": validation, "dataset_version": dataset_version,
            "entity_counts": {name: len(rows) for name, rows in tables.items()},
            "warnings": [issue["message"] for issue in warnings], "source_sha256": source_sha256,
            "source_tables": tables}


class FactoryStateRegistry:
    """SQLite-backed, content-addressed immutable factory-state registry."""

    def __init__(self, db):
        self.db = db
        with db.transaction():
            db.conn.execute("""CREATE TABLE IF NOT EXISTS factory_states (
                state_id TEXT PRIMARY KEY, state_json TEXT NOT NULL,
                validation_json TEXT NOT NULL, metadata_json TEXT NOT NULL,
                source_json TEXT NOT NULL, created_at TEXT NOT NULL)""")

    def load_workbook(self, path: str | Path, as_of_time: str | None = None) -> dict:
        result = normalize_workbook(path, as_of_time)
        identity = {"dataset_version": result["dataset_version"], "state": result["state"],
                    "validation": result["validation"]}
        state_id = hashlib.sha256(canonical(identity).encode()).hexdigest()
        metadata = {"dataset_version": result["dataset_version"], "entity_counts": result["entity_counts"],
                    "warnings": result["warnings"], "source_sha256": result["source_sha256"]}
        values = (state_id, canonical(result["state"]), canonical(result["validation"]), canonical(metadata), canonical(result["source_tables"]), self.db.clock.now())
        with self.db.transaction():
            existing = self.db.conn.execute("SELECT state_json,validation_json,metadata_json,source_json FROM factory_states WHERE state_id=?", (state_id,)).fetchone()
            if existing and tuple(existing) != values[1:5]:
                raise RuntimeError("immutable factory state collision")
            if not existing:
                self.db.conn.execute("INSERT INTO factory_states VALUES(?,?,?,?,?,?)", values)
        return {"state_id": state_id, "dataset_version": metadata["dataset_version"],
                "entity_counts": metadata["entity_counts"], "warnings": metadata["warnings"]}

    def register_normalized(
        self, state: dict, dataset_version: str = "V1.8", *, validation: dict | None = None,
    ) -> dict:
        """Register already-normalized JSON for the same reference-only runtime path."""
        from .runtime_planning import _state_id
        from .v18_adapter import validate_generation_state
        state = deepcopy(state)
        validate_generation_state(state)
        validation = deepcopy(validation) if validation is not None else {
            "status": "VALID", "errors": [], "warnings": [], "quarantined_entities": []
        }
        validate_tool_payload(
            validation, "validate_factory_state", "output_schema", "validate_factory_state_output"
        )
        base_state_id = _state_id(state)
        state_id = base_state_id if validation["status"] == "VALID" else hashlib.sha256(
            canonical({"base_state_id": base_state_id, "validation": validation}).encode("utf-8")
        ).hexdigest()
        counts = {"Orders": len(state["orders"]), "Products": len(state["products"]),
                  "Machines": len(state["machines"]), "Workers": len(state["workers"]),
                  "Shift Calendar": len(state["shifts"]), "Changeovers": len(state["changeovers"])}
        metadata = {"dataset_version": dataset_version, "entity_counts": counts,
                    "warnings": [row["message"] for row in validation["warnings"]],
                    "source_sha256": hashlib.sha256(canonical(state).encode()).hexdigest()}
        values = (state_id, canonical(state), canonical(validation), canonical(metadata), canonical({}), self.db.clock.now())
        with self.db.transaction():
            existing = self.db.conn.execute("SELECT state_json,validation_json,metadata_json,source_json FROM factory_states WHERE state_id=?", (state_id,)).fetchone()
            if existing and tuple(existing) != values[1:5]:
                raise RuntimeError("immutable factory state collision")
            if not existing:
                self.db.conn.execute("INSERT INTO factory_states VALUES(?,?,?,?,?,?)", values)
        return {"state_id": state_id, "dataset_version": dataset_version,
                "entity_counts": counts, "warnings": metadata["warnings"]}

    def get(self, state_id: str) -> dict:
        with self.db.lock:
            row = self.db.conn.execute("SELECT * FROM factory_states WHERE state_id=?", (state_id,)).fetchone()
        if not row:
            raise KeyError("factory state not found")
        return {"state_id": state_id, "state": json.loads(row["state_json"]),
                "validation": json.loads(row["validation_json"]),
                "metadata": json.loads(row["metadata_json"]), "source_tables": json.loads(row["source_json"])}

    def validate(self, state_id: str) -> dict:
        return self.get(state_id)["validation"]

    def planning_state(self, state_id: str) -> dict:
        record = self.get(state_id)
        if record["validation"]["status"] == "INVALID":
            raise ValueError("factory state is structurally INVALID")
        return deepcopy(record["state"])
